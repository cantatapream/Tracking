#!/usr/bin/env python3
"""Build 초과근무수당_집계.xlsm -- a macro-enabled workbook whose button
runs the OvertimeReport VBA module."""
import os
import re
import shutil
import sys
import tempfile
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

import vbaproject

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BAS = os.path.join(ROOT, "OvertimeReport.bas")
OUT = os.path.join(ROOT, "초과근무수당_집계.xlsm")

SHEET_RESULT = "집계결과"
SHEET_LOG = "처리내역"

DRAWING_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"\
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<xdr:twoCellAnchor editAs="oneCell">
<xdr:from><xdr:col>0</xdr:col><xdr:colOff>57150</xdr:colOff>\
<xdr:row>1</xdr:row><xdr:rowOff>38100</xdr:rowOff></xdr:from>
<xdr:to><xdr:col>2</xdr:col><xdr:colOff>228600</xdr:colOff>\
<xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:sp macro="[0]!RunAggregate" textlink="">
<xdr:nvSpPr><xdr:cNvPr id="2" name="btnRunAggregate"/>
<xdr:cNvSpPr><a:spLocks noChangeArrowheads="1"/></xdr:cNvSpPr></xdr:nvSpPr>
<xdr:spPr bwMode="auto">
<a:xfrm><a:off x="57150" y="219075"/><a:ext cx="2124075" cy="466725"/></a:xfrm>
<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 12000"/></a:avLst></a:prstGeom>
<a:solidFill><a:srgbClr val="1F4E79"/></a:solidFill>
<a:ln w="9525"><a:solidFill><a:srgbClr val="17385F"/></a:solidFill></a:ln>
</xdr:spPr>
<xdr:txBody>
<a:bodyPr vertOverflow="clip" horzOverflow="clip" wrap="square" anchor="ctr"/>
<a:lstStyle/>
<a:p><a:pPr algn="ctr"/><a:r>
<a:rPr lang="ko-KR" altLang="en-US" sz="1100" b="1">
<a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
<a:latin typeface="맑은 고딕"/><a:ea typeface="맑은 고딕"/></a:rPr>
<a:t>엑셀 파일 선택 → 집계</a:t></a:r></a:p>
</xdr:txBody>
</xdr:sp>
<xdr:clientData fPrintsWithSheet="0"/>
</xdr:twoCellAnchor>
</xdr:wsDr>"""

SHEET_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\
<Relationship Id="rIdDrw1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"\
 Target="../drawings/drawing1.xml"/></Relationships>"""


def make_base_xlsx(path):
    wb = Workbook()
    wb.code_name = "ThisWorkbook"

    ws = wb.active
    ws.title = SHEET_RESULT
    ws.sheet_properties.codeName = "Sheet1"
    ws["A1"] = "초과근무수당 집계"
    ws["A1"].font = Font(name="맑은 고딕", size=16, bold=True, color="1F4E79")
    ws["A4"] = "아래 버튼을 눌러 월별 초과근무수당 산정서 파일을 선택하세요. (여러 개 선택 가능)"
    ws["A4"].font = Font(name="맑은 고딕", size=9, color="5A5A5A")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    for col, width in (("A", 10), ("B", 10), ("C", 16), ("D", 16)):
        ws.column_dimensions[col].width = width

    wl = wb.create_sheet(SHEET_LOG)
    wl.sheet_properties.codeName = "Sheet2"
    wl["A1"] = "집계를 실행하면 파일별 처리 결과가 여기에 기록됩니다."
    wl["A1"].font = Font(name="맑은 고딕", size=9, color="5A5A5A")
    wl.column_dimensions["A"].width = 42

    wb.save(path)


def repack(src_xlsx, vba_bin, dst_xlsm):
    zin = zipfile.ZipFile(src_xlsx)
    names = zin.namelist()
    if "xl/worksheets/sheet1.xml" not in names:
        raise SystemExit("unexpected package layout: %s" % names)

    os.makedirs(os.path.dirname(dst_xlsm), exist_ok=True)
    zout = zipfile.ZipFile(dst_xlsm, "w", zipfile.ZIP_DEFLATED)

    for item in zin.infolist():
        name = item.filename
        data = zin.read(name)
        if name == "[Content_Types].xml":
            text = data.decode("utf-8")
            text = text.replace(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            if 'Extension="bin"' not in text:
                text = text.replace(
                    "<Types ", '<Types ', 1)
                text = text.replace(
                    "</Types>",
                    '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'
                    '<Override PartName="/xl/drawings/drawing1.xml"'
                    ' ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
                    "</Types>")
            data = text.encode("utf-8")
        elif name == "xl/_rels/workbook.xml.rels":
            text = data.decode("utf-8").replace(
                "</Relationships>",
                '<Relationship Id="rIdVBA1"'
                ' Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject"'
                ' Target="vbaProject.bin"/></Relationships>')
            data = text.encode("utf-8")
        elif name == "xl/worksheets/sheet1.xml":
            text = data.decode("utf-8")
            if "<drawing " in text:
                raise SystemExit("sheet1 already carries a drawing")
            text = text.replace("</worksheet>", '<drawing r:id="rIdDrw1"/></worksheet>')
            if 'xmlns:r=' not in text.split(">", 2)[1]:
                text = re.sub(
                    r"<worksheet ",
                    '<worksheet xmlns:r="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships" ', text, count=1)
            data = text.encode("utf-8")
        elif name == "xl/worksheets/_rels/sheet1.xml.rels":
            raise SystemExit("sheet1 already has relationships; merge needed")
        zout.writestr(item, data)

    zout.writestr("xl/worksheets/_rels/sheet1.xml.rels", SHEET_RELS)
    zout.writestr("xl/drawings/drawing1.xml", DRAWING_XML)
    zout.writestr("xl/vbaProject.bin", vba_bin)
    zout.close()
    zin.close()


def main():
    source = open(BAS, encoding="utf-8").read()
    modules = [
        ("ThisWorkbook", vbaproject.document_module_source("ThisWorkbook", True), True),
        ("Sheet1", vbaproject.document_module_source("Sheet1", False), True),
        ("Sheet2", vbaproject.document_module_source("Sheet2", False), True),
        ("OvertimeReport", source, False),
    ]
    vba_bin = vbaproject.build("VBAProject", modules)

    tmp = tempfile.mkdtemp()
    try:
        base = os.path.join(tmp, "base.xlsx")
        make_base_xlsx(base)
        repack(base, vba_bin, OUT)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # a cp949 copy of the module, for manual import into the VBA editor
    bas_out = os.path.join(ROOT, "OvertimeReport_import.bas")
    with open(bas_out, "w", encoding="cp949", newline="\r\n") as fh:
        fh.write(source)

    print("%s  (%d bytes, vbaProject.bin %d bytes)"
          % (OUT, os.path.getsize(OUT), len(vba_bin)))
    print("%s  (cp949)" % bas_out)


if __name__ == "__main__":
    main()
