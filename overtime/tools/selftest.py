#!/usr/bin/env python3
"""Build sample 산정서 workbooks shaped like the real ones, then run a Python
port of the VBA scan over them to check the anchor logic.

This exercises the algorithm (where the labels and values sit relative to the
[급여 세부내용] anchor, and that a month's day-count does not matter).
It does not execute the VBA itself.
"""
import calendar
import os

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLE_DIR = os.path.join(os.path.dirname(HERE), "sample")

DAY_COL0 = 6          # day columns start at F, as in the real sheet
BLOCK_H = 8           # rows per person


def _person_block(ws, row, seq, name, box_col, days, total_work, ot_raw, ot_paid):
    """One person: the daily grid on the left, the 급여 세부내용 box on the right."""
    ws.cell(row, 1, seq)
    ws.cell(row, 3, name)
    labels = ["구분", "근무시간(시간)", "복무규정시간(시간)", "시간외(시간)",
              "야간(시간)", "휴일(일수)", "출동가산(일수)", "교대시간(분)"]
    for i, lab in enumerate(labels):
        ws.cell(row + i, 5, lab)
    sum_col = DAY_COL0 + days
    for d in range(days):                      # plausible daily numbers
        ws.cell(row + 1, DAY_COL0 + d, 8)
    ws.cell(row + 1, sum_col, total_work)      # 근무시간 합계
    ws.cell(row + 3, sum_col, ot_raw)          # 시간외 합계 (실적 원값)

    # 급여 세부내용 박스
    ws.cell(row, box_col, "급여 세부내용")
    ws.merge_cells(start_row=row, start_column=box_col, end_row=row, end_column=box_col + 1)
    ws.cell(row + 1, box_col, "총 근무시간")
    ws.cell(row + 1, box_col + 1, "총 복무규정시간")
    ws.cell(row + 2, box_col, total_work)
    ws.cell(row + 2, box_col + 1, 120)
    ws.cell(row + 3, box_col, "시간외(시간)")
    ws.cell(row + 3, box_col + 1, ot_paid)     # 상한이 적용된 지급 기준 값
    ws.cell(row + 4, box_col, "야간(시간)")
    ws.cell(row + 4, box_col + 1, 56)
    ws.cell(row + 5, box_col, "휴일(일수)")
    ws.cell(row + 5, box_col + 1, 2)
    ws.cell(row + 6, box_col, "출동가산(일수)")
    ws.cell(row + 6, box_col + 1, 8)


def make_sample(path, year, month, ship, people, sheet_name, decoy=False):
    days = calendar.monthrange(year, month)[1]
    box_col = DAY_COL0 + days + 3
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    ws.cell(1, 1, "초과근무수당 개인별 산출내역서(%s)" % ship)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws.cell(9, 1, "초과근무수당 일집계 현황(%d년 %d월)" % (year, month))
    ws.cell(10, 1, "순번"); ws.cell(10, 3, "성명"); ws.cell(10, 5, "일자")
    for d in range(days):
        ws.cell(10, DAY_COL0 + d, d + 1)
    ws.cell(10, DAY_COL0 + days, "합계")

    row = 11
    for seq, (name, work, ot_raw, ot_paid) in enumerate(people, start=1):
        _person_block(ws, row, seq, name, box_col, days, work, ot_raw, ot_paid)
        row += BLOCK_H

    if decoy:
        # a second sheet carrying the same title -- the sheet-name rule must skip it
        w2 = wb.create_sheet("%d월 초과 합계 내역" % month)
        w2.cell(1, 1, "초과근무수당 개인별 산출내역서(%s)" % ship)
        w2.cell(9, 1, "초과근무수당 일집계 현황(%d년 %d월)" % (year, month))
        _person_block(w2, 11, 1, "중복", box_col, days, 9999, 9999, 9999)

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# Python port of the VBA scan (same anchors, same relative offsets)
# --------------------------------------------------------------------------
K_TITLE, K_PERIOD, K_BOX = "초과근무수당개인별산출내역서", "일집계현황", "급여세부내용"
K_WORK, K_OT = "총근무시간", "시간외"


def norm(v):
    if not isinstance(v, str):
        return ""
    for ch in (" ", "\t", "\r", "\n", " ", "　"):
        v = v.replace(ch, "")
    return v


def tail_digits(s):
    t = ""
    for ch in reversed(s):
        if ch.isdigit():
            t = ch + t
        else:
            break
    return int(t) if 0 < len(t) <= 4 else 0


def parse_ym(s):
    p, q = s.find("년"), s.find("월")
    if p < 0 or q <= p:
        return 0, 0
    y, m = tail_digits(s[:p]), tail_digits(s[p + 1:q])
    if 0 <= y <= 99:
        y += 2000
    if not (1990 <= y <= 2099) or not (1 <= m <= 12):
        return 0, 0
    return y, m


def try_num(grid, r, c):
    if r < 1 or c < 1 or r > len(grid) or c > len(grid[0]):
        return None
    v = grid[r - 1][c - 1]
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = norm(v).replace(",", "")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def pick_num(grid, r1, c1, r2, c2):
    for r, c in ((r1, c1), (r2, c2)):
        d = try_num(grid, r, c)
        if d is not None:
            return d
    return 0.0


def sum_box(grid, r0, c0):
    work = ot = 0.0
    got_w = got_o = False
    for r in range(r0 + 1, min(r0 + 10, len(grid)) + 1):
        for c in range(c0, min(c0 + 2, len(grid[0])) + 1):
            s = norm(grid[r - 1][c - 1])
            if not got_w and s.startswith(K_WORK):
                work += pick_num(grid, r + 1, c, r, c + 1)
                got_w = True
            if not got_o and s.startswith(K_OT):
                ot += pick_num(grid, r, c + 1, r + 1, c)
                got_o = True
        if got_w and got_o:
            break
    return work, ot


def sheet_has_title(grid):
    for row in grid[:25]:
        for v in row[:120]:
            if K_TITLE in norm(v):
                return True
    return False


def pick_sheet_by_name(names):
    hits = [n for n in names
            if ("초과근무수당" in norm(n) or "초간근무수당" in norm(n))
            and "합계" not in norm(n) and "내역" not in norm(n)]
    return hits[0] if len(hits) == 1 else None


def scan_file(path):
    wb = load_workbook(path, data_only=True, read_only=False)
    grids = {}
    for ws in wb.worksheets:
        grid = [[c.value for c in row] for row in ws.iter_rows()]
        if grid and sheet_has_title(grid):
            grids[ws.title] = grid
    if not grids:
        return {"status": "산출내역서 시트 없음"}
    if len(grids) == 1:
        name = next(iter(grids))
    else:
        name = pick_sheet_by_name(list(grids))
        if name is None:
            return {"status": "대상 시트 판별 실패"}
    grid = grids[name]

    year = month = 0
    ship = ""
    for row in grid[:25]:
        for v in row:
            s = norm(v)
            if not s:
                continue
            if year == 0 and K_PERIOD in s:
                year, month = parse_ym(s)
            if not ship and K_TITLE in s and "(" in s and ")" in s:
                ship = s[s.rindex("(") + 1:s.index(")", s.rindex("("))]

    work = ot = 0.0
    people = 0
    for r in range(1, len(grid) + 1):
        for c in range(1, len(grid[0]) + 1):
            if K_BOX in norm(grid[r - 1][c - 1]):
                people += 1
                w, o = sum_box(grid, r, c)
                work += w
                ot += o
    return {"status": "정상", "sheet": name, "year": year, "month": month,
            "ship": ship, "people": people, "work": work, "ot": ot}


def main():
    # 현관용(3번)처럼 시간외 실적(108)과 지급 기준(100)이 다른 사람을 포함시킨다
    may = [("김근홍", 220, 100, 100), ("최영필", 232, 100, 100), ("현관용", 208, 108, 100)]
    feb = [("김근홍", 190, 96, 96), ("최영필", 205, 130, 100)]

    f5 = make_sample(os.path.join(SAMPLE_DIR, "5월 초과근무수당 산정서(1505함).xlsx"),
                     2025, 5, "1505함", may, "5월 초간근무수당", decoy=True)
    f2 = make_sample(os.path.join(SAMPLE_DIR, "2월 초과근무수당 산정서(1505함).xlsx"),
                     2024, 2, "1505함", feb, "2월 초과근무수당")

    ok = True
    for path, exp in ((f5, dict(year=2025, month=5, people=3, work=660.0, ot=300.0)),
                      (f2, dict(year=2024, month=2, people=2, work=395.0, ot=196.0))):
        got = scan_file(path)
        print(os.path.basename(path))
        print("   ", got)
        for k, v in exp.items():
            if got.get(k) != v:
                print("    MISMATCH %s: got %r want %r" % (k, got.get(k), v))
                ok = False
    print()
    print("sample files ->", SAMPLE_DIR)
    print("RESULT:", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
